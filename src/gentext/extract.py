"""Source file → plain Markdown-ish text (DR-0011 T1 input).

Deterministic, no LLM. Output is cached under build/text/ (derived, gitignored).
"""

from __future__ import annotations

import re
import subprocess
from pathlib import Path

BUILD_TEXT = Path("build/text")

_COMMENT_JSON_MARKER = "<!-- connector commentThreads"
_KIX = re.compile(r"</?comment_(?:start|end)[^>]*>")
_FRONTMATTER = re.compile(r"\A---\n.*?\n---\n", re.S)


def _from_md(path: Path) -> str:
    text = path.read_text(errors="replace")
    text = _FRONTMATTER.sub("", text, count=1)
    cut = text.find(_COMMENT_JSON_MARKER)
    if cut != -1:
        text = text[:cut]
    return _KIX.sub("", text)


def _from_pdf(path: Path) -> str:
    out = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, check=True)
    return out.stdout


def _from_docx(path: Path) -> str:
    import pypandoc

    return pypandoc.convert_file(str(path), "gfm", format="docx", extra_args=["--wrap=none"])


def _from_xlsx(path: Path, max_rows: int = 400) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append("… (more rows truncated)\n")
                break
            cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
            if any(cells):
                parts.append("\t".join(cells) + "\n")
    return "".join(parts)


EXTRACTORS = {".md": _from_md, ".pdf": _from_pdf, ".docx": _from_docx, ".xlsx": _from_xlsx}


def extract(path: Path) -> str:
    fn = EXTRACTORS.get(path.suffix.lower())
    if fn is None:
        raise ValueError(f"no extractor for {path.suffix}: {path}")
    return fn(path)


def cached_text(asset_id: str, path: Path) -> str:
    """Extract once per source file; re-extract if the source is newer than the cache."""
    BUILD_TEXT.mkdir(parents=True, exist_ok=True)
    cache = BUILD_TEXT / f"{asset_id}.md"
    if cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        return cache.read_text()
    text = extract(path)
    cache.write_text(text)
    return text


EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?<!\d)(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}(?!\d)")


def redact(text: str) -> str:
    """Remove personal contact details before text leaves the machine (DR-0004)."""
    return PHONE.sub("[phone]", EMAIL.sub("[email]", text))
